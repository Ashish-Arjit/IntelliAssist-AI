"""Document loader module for extracting text and metadata from various file formats.

Supports PDF, TXT, and DOCX files with rich metadata extraction.
"""

from io import BytesIO
import os
from typing import BinaryIO, List, Optional, Union
import docx
from langchain_core.documents import Document
from pypdf import PdfReader


class DocumentLoader:
    """Modular document loader for processing uploaded or local files."""

    @staticmethod
    def _get_stream_and_filename(
        file_source: Union[str, bytes, BinaryIO], default_name: str
    ) -> tuple[BytesIO, str]:
        """Extract a BytesIO stream and resolved filename from various input types."""
        filename = default_name
        if isinstance(file_source, str):
            if not os.path.exists(file_source):
                raise FileNotFoundError(f"File not found: {file_source}")
            with open(file_source, "rb") as f:
                content = f.read()
            stream = BytesIO(content)
            filename = os.path.basename(file_source)
        elif isinstance(file_source, bytes):
            stream = BytesIO(file_source)
        else:
            if hasattr(file_source, "getvalue"):
                stream = BytesIO(file_source.getvalue())
            else:
                stream = BytesIO(file_source.read())
            if hasattr(file_source, "name") and file_source.name:
                filename = file_source.name

        return stream, filename

    @classmethod
    def load_pdf(
        cls, file_source: Union[str, bytes, BinaryIO], filename: Optional[str] = None
    ) -> List[Document]:
        """Extract text and metadata from a PDF file per page.

        Args:
            file_source: File path, bytes, or binary stream.
            filename: Optional display name for the document.

        Returns:
            List of LangChain Document objects with page content and metadata.
        """
        stream, resolved_filename = cls._get_stream_and_filename(
            file_source, filename or "document.pdf"
        )
        if filename:
            resolved_filename = filename

        reader = PdfReader(stream)
        total_pages = len(reader.pages)
        documents: List[Document] = []

        for page_idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            metadata = {
                "source": resolved_filename,
                "file_name": resolved_filename,
                "file_type": "pdf",
                "page": page_idx,
                "total_pages": total_pages,
                "character_count": len(text),
            }
            documents.append(Document(page_content=text, metadata=metadata))

        return documents

    @classmethod
    def load_txt(
        cls,
        file_source: Union[str, bytes, BinaryIO],
        filename: Optional[str] = None,
        encoding: str = "utf-8",
    ) -> List[Document]:
        """Extract text and metadata from a plain text (TXT) document.

        Args:
            file_source: File path, bytes, or binary stream.
            filename: Optional display name for the document.
            encoding: Text encoding (defaults to utf-8).

        Returns:
            List containing a single LangChain Document object.
        """
        stream, resolved_filename = cls._get_stream_and_filename(
            file_source, filename or "document.txt"
        )
        if filename:
            resolved_filename = filename

        raw_bytes = stream.getvalue()
        try:
            text = raw_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            text = raw_bytes.decode("latin-1", errors="replace")

        metadata = {
            "source": resolved_filename,
            "file_name": resolved_filename,
            "file_type": "txt",
            "character_count": len(text),
            "line_count": len(text.splitlines()),
        }

        return [Document(page_content=text, metadata=metadata)]

    @classmethod
    def load_docx(
        cls, file_source: Union[str, bytes, BinaryIO], filename: Optional[str] = None
    ) -> List[Document]:
        """Extract text and metadata from a Microsoft Word (DOCX) document.

        Args:
            file_source: File path, bytes, or binary stream.
            filename: Optional display name for the document.

        Returns:
            List containing a LangChain Document object.
        """
        stream, resolved_filename = cls._get_stream_and_filename(
            file_source, filename or "document.docx"
        )
        if filename:
            resolved_filename = filename

        doc = docx.Document(stream)
        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text if present
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_text.append(row_text)

        all_text_parts = paragraphs_text + table_text
        full_text = "\n\n".join(all_text_parts)

        metadata = {
            "source": resolved_filename,
            "file_name": resolved_filename,
            "file_type": "docx",
            "character_count": len(full_text),
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        }

        return [Document(page_content=full_text, metadata=metadata)]

    @classmethod
    def load_document(
        cls,
        file_source: Union[str, bytes, BinaryIO],
        filename: Optional[str] = None,
        file_type: Optional[str] = None,
    ) -> List[Document]:
        """Unified document loader routing based on detected or specified file type.

        Args:
            file_source: File path, bytes, or binary stream.
            filename: Name of the file.
            file_type: Explicit file extension type ('pdf', 'txt', 'docx').

        Returns:
            List of LangChain Document objects.
        """
        resolved_name = filename
        if not resolved_name:
            if isinstance(file_source, str):
                resolved_name = os.path.basename(file_source)
            elif hasattr(file_source, "name") and file_source.name:
                resolved_name = file_source.name

        inferred_ext = ""
        if resolved_name:
            inferred_ext = resolved_name.rsplit(".", 1)[-1].lower() if "." in resolved_name else ""

        target_type = (file_type or inferred_ext).lower()

        if target_type == "pdf":
            return cls.load_pdf(file_source, filename=resolved_name)
        elif target_type == "txt":
            return cls.load_txt(file_source, filename=resolved_name)
        elif target_type in ("docx", "doc"):
            return cls.load_docx(file_source, filename=resolved_name)
        else:
            raise ValueError(
                f"Unsupported document format: '{target_type}'. "
                "Supported formats are PDF, TXT, and DOCX."
            )

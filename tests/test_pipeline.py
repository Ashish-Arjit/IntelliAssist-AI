"""End-to-end pipeline integration tests for IntelliAssist AI."""

from io import BytesIO
import unittest
import docx
from pypdf import PageObject, PdfWriter

from modules.chunker import DocumentChunker
from modules.document_loader import DocumentLoader
from modules.embeddings import EmbeddingManager
from modules.preprocessor import TextPreprocessor
from modules.validator import DocumentValidator
from modules.vector_store import VectorStoreManager


class TestPipelineIntegration(unittest.TestCase):
    """Integration test suite for the complete document ingestion, embedding, and search pipeline."""

    @classmethod
    def setUpClass(cls):
        cls.embedding_manager = EmbeddingManager()

    def test_full_pipeline_txt(self):
        raw_text = "IntelliAssist AI is a modular smart assistant.\n\n\n\nIt processes documents efficiently."
        txt_bytes = raw_text.encode("utf-8")

        # 1. Validation
        val = DocumentValidator.validate_file_upload(
            type("MockFile", (), {"name": "doc.txt", "size": len(txt_bytes), "getvalue": lambda self: txt_bytes})()
        )
        self.assertTrue(val.is_valid)

        # 2. Loading
        docs = DocumentLoader.load_document(txt_bytes, filename="doc.txt")
        self.assertEqual(len(docs), 1)

        # 3. Content Validation
        c_val = DocumentValidator.validate_extracted_content(docs, filename="doc.txt")
        self.assertTrue(c_val.is_valid)

        # 4. Preprocessing
        preprocessed = TextPreprocessor.preprocess_documents(docs)
        self.assertEqual(len(preprocessed), 1)
        self.assertEqual(
            preprocessed[0].page_content,
            "IntelliAssist AI is a modular smart assistant.\n\nIt processes documents efficiently.",
        )

        # 5. Chunking
        chunker = DocumentChunker(chunk_size=50, chunk_overlap=10)
        chunks = chunker.split_documents(preprocessed)
        self.assertGreater(len(chunks), 1)
        self.assertEqual(chunks[0].metadata["file_name"], "doc.txt")
        self.assertEqual(chunks[0].metadata["chunk_index"], 1)

        # 6. FAISS Embedding and Search
        vector_manager = VectorStoreManager(embedding_manager=self.embedding_manager)
        vector_manager.create_from_documents(chunks)
        self.assertEqual(vector_manager.total_vectors, len(chunks))

        results = vector_manager.similarity_search_with_score("modular smart assistant", top_k=2)
        self.assertGreater(len(results), 0)
        self.assertIn("IntelliAssist", results[0][0].page_content)
        self.assertGreater(results[0][2], 0.4)

    def test_full_pipeline_docx(self):
        doc = docx.Document()
        doc.add_heading("Machine Learning Architecture", level=1)
        doc.add_paragraph("Neural networks learn latent representations from data through backpropagation.")
        stream = BytesIO()
        doc.save(stream)
        docx_bytes = stream.getvalue()

        docs = DocumentLoader.load_document(docx_bytes, filename="ml_spec.docx")
        preprocessed = TextPreprocessor.preprocess_documents(docs)
        chunker = DocumentChunker(chunk_size=100, chunk_overlap=20)
        chunks = chunker.split_documents(preprocessed)

        self.assertGreaterEqual(len(chunks), 1)
        self.assertEqual(chunks[0].metadata["file_type"], "docx")

        vector_manager = VectorStoreManager(embedding_manager=self.embedding_manager)
        vector_manager.create_from_documents(chunks)
        results = vector_manager.similarity_search_with_score("deep neural networks backpropagation", top_k=1)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0][0].metadata["file_name"], "ml_spec.docx")

    def test_full_pipeline_pdf(self):
        from tests.test_vector_search_integration import generate_test_pdf_bytes

        pdf_bytes = generate_test_pdf_bytes([
            "IntelliAssist AI supports document processing for PDF files.\nPreserving page numbers and metadata is essential.",
            "Second page details: Semantic search utilizes vector cosine similarity."
        ])

        # 1. Validation
        val = DocumentValidator.validate_file_upload(
            type("MockFile", (), {"name": "manual.pdf", "size": len(pdf_bytes), "getvalue": lambda self: pdf_bytes})()
        )
        self.assertTrue(val.is_valid)

        # 2. Loading
        docs = DocumentLoader.load_document(pdf_bytes, filename="manual.pdf")
        self.assertEqual(len(docs), 2)
        self.assertEqual(docs[0].metadata["file_type"], "pdf")
        self.assertEqual(docs[0].metadata["page"], 1)
        self.assertEqual(docs[1].metadata["page"], 2)

        # 3. Preprocessing
        preprocessed = TextPreprocessor.preprocess_documents(docs)
        self.assertEqual(len(preprocessed), 2)

        # 4. Chunking
        chunker = DocumentChunker(chunk_size=120, chunk_overlap=20)
        chunks = chunker.split_documents(preprocessed)
        self.assertGreaterEqual(len(chunks), 2)
        self.assertEqual(chunks[0].metadata["file_name"], "manual.pdf")

        # 5. Vector search
        vector_manager = VectorStoreManager(embedding_manager=self.embedding_manager)
        vector_manager.create_from_documents(chunks)
        self.assertEqual(vector_manager.total_vectors, len(chunks))

        results = vector_manager.similarity_search_with_score("vector cosine similarity", top_k=1)
        self.assertGreater(len(results), 0)
        self.assertEqual(results[0][0].metadata["file_name"], "manual.pdf")
        self.assertEqual(results[0][0].metadata["page"], 2)


if __name__ == "__main__":
    unittest.main()

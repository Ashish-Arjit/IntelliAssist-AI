"""Unit tests for DocumentLoader (PDF, TXT, DOCX, and dispatching)."""

from io import BytesIO
import unittest
import docx
from pypdf import PageObject, PdfWriter
from modules.document_loader import DocumentLoader


class TestDocumentLoader(unittest.TestCase):
    """Test suite for DocumentLoader across PDF, TXT, and DOCX formats."""

    def test_load_pdf(self):
        writer = PdfWriter()
        page1 = PageObject.create_blank_page(width=200, height=200)
        page2 = PageObject.create_blank_page(width=200, height=200)
        writer.add_page(page1)
        writer.add_page(page2)
        stream = BytesIO()
        writer.write(stream)

        docs = DocumentLoader.load_pdf(stream.getvalue(), filename="report.pdf")
        self.assertEqual(len(docs), 2)
        self.assertEqual(docs[0].metadata["file_name"], "report.pdf")
        self.assertEqual(docs[0].metadata["file_type"], "pdf")
        self.assertEqual(docs[0].metadata["page"], 1)
        self.assertEqual(docs[0].metadata["total_pages"], 2)
        self.assertEqual(docs[1].metadata["page"], 2)

    def test_load_txt(self):
        sample_text = "Line 1: Hello IntelliAssist\nLine 2: Smart Assistant\nLine 3: End"
        txt_bytes = sample_text.encode("utf-8")

        docs = DocumentLoader.load_txt(txt_bytes, filename="notes.txt")
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].page_content, sample_text)
        self.assertEqual(docs[0].metadata["file_name"], "notes.txt")
        self.assertEqual(docs[0].metadata["file_type"], "txt")
        self.assertEqual(docs[0].metadata["line_count"], 3)
        self.assertEqual(docs[0].metadata["character_count"], len(sample_text))

    def test_load_docx(self):
        doc = docx.Document()
        doc.add_heading("Test Document", level=1)
        doc.add_paragraph("Paragraph 1 content.")
        doc.add_paragraph("Paragraph 2 content.")
        
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(0, 1).text = "B1"
        table.cell(1, 0).text = "A2"
        table.cell(1, 1).text = "B2"

        stream = BytesIO()
        doc.save(stream)
        docx_bytes = stream.getvalue()

        docs = DocumentLoader.load_docx(docx_bytes, filename="overview.docx")
        self.assertEqual(len(docs), 1)
        self.assertIn("Test Document", docs[0].page_content)
        self.assertIn("Paragraph 1 content.", docs[0].page_content)
        self.assertIn("A1 | B1", docs[0].page_content)
        self.assertEqual(docs[0].metadata["file_name"], "overview.docx")
        self.assertEqual(docs[0].metadata["file_type"], "docx")
        self.assertGreaterEqual(docs[0].metadata["paragraph_count"], 3)
        self.assertEqual(docs[0].metadata["table_count"], 1)

    def test_load_document_dispatch(self):
        txt_bytes = b"Sample content"
        docs = DocumentLoader.load_document(txt_bytes, filename="data.txt")
        self.assertEqual(docs[0].metadata["file_type"], "txt")

        with self.assertRaises(ValueError):
            DocumentLoader.load_document(b"fake", filename="data.xyz")


if __name__ == "__main__":
    unittest.main()

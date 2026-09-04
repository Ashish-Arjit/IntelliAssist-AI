"""Integration tests for multi-document ingestion and semantic vector search across PDF, TXT, and DOCX."""

from io import BytesIO
import unittest
import docx

from modules.chunker import DocumentChunker
from modules.document_loader import DocumentLoader
from modules.embeddings import EmbeddingManager
from modules.preprocessor import TextPreprocessor
from modules.validator import DocumentValidator
from modules.vector_store import VectorStoreManager


def generate_test_pdf_bytes(pages: list[str]) -> bytes:
    """Generate a clean in-memory PDF with readable text per page for testing."""
    page_obj_ids = []
    content_obj_ids = []
    next_id = 3
    for _ in pages:
        page_obj_ids.append(next_id)
        content_obj_ids.append(next_id + 1)
        next_id += 2
    font_id = next_id

    objects = []
    objects.append(b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n")
    kids_str = " ".join(f"{pid} 0 R" for pid in page_obj_ids)
    objects.append(f"2 0 obj\n<< /Type /Pages /Kids [{kids_str}] /Count {len(pages)} >>\nendobj\n".encode("latin-1"))

    for pid, cid, text in zip(page_obj_ids, content_obj_ids, pages):
        p_obj = (
            f"{pid} 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {cid} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>\nendobj\n"
        ).encode("latin-1")
        objects.append(p_obj)
        escaped_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_data = f"BT /F1 12 Tf 72 712 Td ({escaped_text}) Tj ET".encode("latin-1")
        c_obj = (
            f"{cid} 0 obj\n<< /Length {len(stream_data)} >>\nstream\n".encode("latin-1")
            + stream_data
            + b"\nendstream\nendobj\n"
        )
        objects.append(c_obj)

    f_obj = f"{font_id} 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n".encode("latin-1")
    objects.append(f_obj)

    header = b"%PDF-1.4\n"
    body = bytearray(header)
    offsets = [0]
    for obj in objects:
        offsets.append(len(body))
        body.extend(obj)
    startxref = len(body)
    xref = f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode("latin-1")
    for off in offsets[1:]:
        xref += f"{off:010d} 00000 n \n".encode("latin-1")
    trailer = f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{startxref}\n%%EOF\n".encode("latin-1")
    body.extend(xref)
    body.extend(trailer)
    return bytes(body)


class TestVectorSearchIntegration(unittest.TestCase):
    """End-to-end multi-format ingestion and vector search integration tests."""

    @classmethod
    def setUpClass(cls):
        cls.embedding_manager = EmbeddingManager()
        cls.chunker = DocumentChunker(chunk_size=400, chunk_overlap=50)

        # 1. Create a 2-page PDF document about Neural Architectures & Machine Learning
        cls.pdf_bytes = generate_test_pdf_bytes([
            "Artificial Intelligence Architecture: Deep learning systems utilize multi-layer perceptrons and backpropagation.",
            "Convolutional and recurrent networks process spatial and sequential patterns respectively in AI research."
        ])

        # 2. Create a TXT document about Cloud Computing and Kubernetes
        cls.txt_content = (
            "Cloud Computing Architecture:\n"
            "Kubernetes is an open-source system for automating deployment, scaling, "
            "and management of containerized applications.\n"
            "Docker containers package software into standardized units for development and shipment."
        )
        cls.txt_bytes = cls.txt_content.encode("utf-8")

        # 3. Create a DOCX document about Financial Quarterly Results
        doc = docx.Document()
        doc.add_heading("Q3 Financial Performance Report", level=1)
        doc.add_paragraph("Revenue reached $12.5 million, representing a 24% year-over-year increase.")
        doc.add_paragraph("Operating expenses decreased by 8% due to streamlined cloud infrastructure.")
        docx_stream = BytesIO()
        doc.save(docx_stream)
        cls.docx_bytes = docx_stream.getvalue()

    def test_multi_document_indexing_and_cross_search(self):
        """Test processing PDF, TXT, and DOCX simultaneously into FAISS and querying across them."""
        # --- 1. Process PDF ---
        val_pdf = DocumentValidator.validate_file_upload(
            type("MockFile", (), {"name": "ai_research.pdf", "size": len(self.pdf_bytes), "getvalue": lambda self: self.pdf_bytes})()
        )
        self.assertTrue(val_pdf.is_valid)
        raw_pdf = DocumentLoader.load_document(self.pdf_bytes, filename="ai_research.pdf")
        self.assertEqual(len(raw_pdf), 2)  # 2 pages
        self.assertEqual(raw_pdf[0].metadata["page"], 1)
        self.assertEqual(raw_pdf[1].metadata["page"], 2)
        prep_pdf = TextPreprocessor.preprocess_documents(raw_pdf)
        chunks_pdf = self.chunker.split_documents(prep_pdf)

        # --- 2. Process TXT ---
        val_txt = DocumentValidator.validate_file_upload(
            type("MockFile", (), {"name": "cloud_guide.txt", "size": len(self.txt_bytes), "getvalue": lambda self: self.txt_bytes})()
        )
        self.assertTrue(val_txt.is_valid)
        raw_txt = DocumentLoader.load_document(self.txt_bytes, filename="cloud_guide.txt")
        prep_txt = TextPreprocessor.preprocess_documents(raw_txt)
        chunks_txt = self.chunker.split_documents(prep_txt)

        # --- 3. Process DOCX ---
        val_docx = DocumentValidator.validate_file_upload(
            type("MockFile", (), {"name": "finance_q3.docx", "size": len(self.docx_bytes), "getvalue": lambda self: self.docx_bytes})()
        )
        self.assertTrue(val_docx.is_valid)
        raw_docx = DocumentLoader.load_document(self.docx_bytes, filename="finance_q3.docx")
        prep_docx = TextPreprocessor.preprocess_documents(raw_docx)
        chunks_docx = self.chunker.split_documents(prep_docx)

        # Combine chunks from all three formats
        all_chunks = chunks_pdf + chunks_txt + chunks_docx
        self.assertGreaterEqual(len(all_chunks), 3)

        # Verify all source formats are represented with proper metadata
        formats = {c.metadata["file_type"] for c in all_chunks}
        self.assertEqual(formats, {"pdf", "txt", "docx"})

        # Index all chunks into FAISS vector store
        vector_manager = VectorStoreManager(embedding_manager=self.embedding_manager)
        vector_manager.create_from_documents(all_chunks)
        self.assertEqual(vector_manager.total_vectors, len(all_chunks))

        # Query 1: Machine Learning & Perceptrons (Should retrieve PDF chunk with page metadata)
        results_ai = vector_manager.similarity_search_with_score("multi-layer perceptrons deep learning", top_k=2)
        self.assertGreater(len(results_ai), 0)
        top_ai_doc, dist_ai, score_ai = results_ai[0]
        self.assertEqual(top_ai_doc.metadata["file_name"], "ai_research.pdf")
        self.assertEqual(top_ai_doc.metadata["file_type"], "pdf")
        self.assertIsNotNone(top_ai_doc.metadata.get("page"))
        self.assertIn("perceptron", top_ai_doc.page_content.lower())
        self.assertGreater(score_ai, 0.4)

        # Query 2: Cloud & Kubernetes (Should retrieve TXT chunk)
        results_cloud = vector_manager.similarity_search_with_score("What is Kubernetes container orchestration?", top_k=2)
        self.assertGreater(len(results_cloud), 0)
        top_cloud_doc, dist_cloud, score_cloud = results_cloud[0]
        self.assertEqual(top_cloud_doc.metadata["file_name"], "cloud_guide.txt")
        self.assertEqual(top_cloud_doc.metadata["file_type"], "txt")
        self.assertIn("Kubernetes", top_cloud_doc.page_content)
        self.assertGreater(score_cloud, 0.4)

        # Query 3: Financial revenue (Should retrieve DOCX chunk)
        results_finance = vector_manager.similarity_search_with_score("What was the quarterly revenue growth percentage?", top_k=2)
        self.assertGreater(len(results_finance), 0)
        top_fin_doc, dist_fin, score_fin = results_finance[0]
        self.assertEqual(top_fin_doc.metadata["file_name"], "finance_q3.docx")
        self.assertEqual(top_fin_doc.metadata["file_type"], "docx")
        self.assertIn("Revenue reached $12.5 million", top_fin_doc.page_content)
        self.assertGreater(score_fin, 0.4)

        # Query 4: Unrelated query should yield significantly lower similarity score
        results_unrelated = vector_manager.similarity_search_with_score("Ancient Roman aqueduct engineering techniques", top_k=1)
        self.assertGreater(len(results_unrelated), 0)
        _, _, score_unrelated = results_unrelated[0]
        self.assertLess(score_unrelated, score_cloud)


if __name__ == "__main__":
    unittest.main()
